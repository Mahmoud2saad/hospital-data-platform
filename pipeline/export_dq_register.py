"""
export_dq_register.py
=====================
Builds a formatted Data Quality & Exception Register (.xlsx) straight from
ctl.dq_exceptions - generated from the pipeline's own exception flags, never
typed up by hand. Reads the LATEST run by default.

"Category" and "Business Impact" are reviewer-facing prose, one sentence per
rule_code, kept here (not in SQL). A rule_code with no entry fails loudly
rather than shipping a register with blank narrative cells.

    python pipeline/run_pipeline.py
    python pipeline/export_dq_register.py
"""
from __future__ import annotations

import argparse
from pathlib import Path

import duckdb
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parent.parent

CATEGORY = {
    "DUP_SOURCE_VERSION": "Duplicate / Versioning",
    "NON_INPATIENT_ENCOUNTER": "Business Rule Exclusion",
    "ORPHAN_DOCTOR_FK": "Referential Integrity",
    "ORPHAN_BILLING_ADMISSION": "Referential Integrity",
    "NEGATIVE_LOS": "Temporal / Date Logic",
    "MISSING_DISCHARGE": "Temporal / Date Logic",
    "SCD_OVERLAP": "Temporal / Date Logic (SCD)",
    "SCD_COVERAGE_GAP": "Temporal / Date Logic (SCD)",
    "SOURCE_DELETED": "Change Data Capture",
    "ILLOGICAL_ADMIT_AFTER_DEATH": "Logical Inconsistency",
    "INSUFFICIENT_OBSERVATION_WINDOW": "Statistical Validity",
}
BUSINESS_IMPACT = {
    "DUP_SOURCE_VERSION": "Double-counts admissions/revenue if not deduplicated; inflates volume and financial KPIs.",
    "NON_INPATIENT_ENCOUNTER": "Inflates admission counts and average charge if test/cancelled encounters are included.",
    "ORPHAN_DOCTOR_FK": "Breaks physician-level attribution and productivity reporting for the affected admission.",
    "ORPHAN_BILLING_ADMISSION": "Revenue that cannot be tied to a clinical encounter; branch/department revenue and "
                                "charge-vs-bill analysis are incomplete until reconciled.",
    "NEGATIVE_LOS": "Corrupts length-of-stay and any date-sequenced KPI (readmission window) if not quarantined.",
    "MISSING_DISCHARGE": "Cannot determine LOS or readmission eligibility; the encounter may still be genuinely open.",
    "SCD_OVERLAP": "Physician attribution is ambiguous for the period; productivity/quality reporting for two org "
                   "units may double-count or misattribute the encounter.",
    "SCD_COVERAGE_GAP": "No authoritative physician assignment for the period; defaulting to \"current\" would "
                        "misattribute historical activity.",
    "SOURCE_DELETED": "A record deleted at source would keep contributing to volume and revenue KPIs if the delete "
                      "event were ignored.",
    "ILLOGICAL_ADMIT_AFTER_DEATH": "Likely patient-identity (MRN reuse/merge) error upstream; uncorrected, it feeds a "
                                   "false readmission and distorts mortality-linked KPIs.",
    "INSUFFICIENT_OBSERVATION_WINDOW": "Readmission rate would be biased downward if these were force-included in the "
                                       "denominator before their 30-day window closes.",
}
NAVY = "1F4E78"
SEV_FILL = {s: PatternFill(start_color=c, end_color=c, fill_type="solid")
            for s, c in (("HIGH", "F8CBAD"), ("MEDIUM", "FFE699"), ("LOW", "C6E0B4"))}


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--db", type=Path, default=ROOT / "output" / "warehouse.duckdb")
    ap.add_argument("--out", type=Path, default=ROOT / "docs" / "data_quality_exception_register.xlsx")
    ap.add_argument("--max-rows", type=int, default=2000, help="cap detail rows (summary always covers all)")
    a = ap.parse_args(argv)
    if not a.db.exists():
        raise SystemExit(f"{a.db} not found - run pipeline/run_pipeline.py first")

    con = duckdb.connect(str(a.db), read_only=True)
    run_id = con.execute("SELECT MAX(run_id) FROM ctl.dq_exceptions").fetchone()[0]
    rows = con.execute("""
        SELECT exception_id, run_id, entity, business_key, rule_code, severity, description, treatment
        FROM ctl.dq_exceptions WHERE run_id = ?
        ORDER BY CASE severity WHEN 'HIGH' THEN 0 WHEN 'MEDIUM' THEN 1 ELSE 2 END, rule_code, business_key""",
                       [run_id]).fetchall()
    summary = con.execute("""
        SELECT rule_code, severity, entity, COUNT(*) FROM ctl.dq_exceptions WHERE run_id = ?
        GROUP BY rule_code, severity, entity ORDER BY CASE severity WHEN 'HIGH' THEN 0 WHEN 'MEDIUM' THEN 1 ELSE 2 END, rule_code""",
                          [run_id]).fetchall()
    con.close()

    unknown = {r[4] for r in rows} - set(CATEGORY)
    if unknown:
        raise SystemExit(f"rule_code(s) with no Category/Impact defined: {sorted(unknown)} - add to this script.")

    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    hfill, hfont = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid"), Font(bold=True, color="FFFFFF")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Meridian Health Group (synthetic data) - Data Quality & Exception Register"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = (f"Generated from ctl.dq_exceptions, run {run_id}. Every row is raised by a pipeline rule, "
                f"not a manual review. Trended across runs via ctl.vw_dq_trend.")
    ws["A2"].font = Font(italic=True, size=9, color="666666")
    for j, h in enumerate(["Rule", "Severity", "Entity", "Category", "Count", "Business impact"], 1):
        c = ws.cell(row=4, column=j, value=h)
        c.font, c.fill, c.border = hfont, hfill, border
    for i, (rule, sev, ent, n) in enumerate(summary, start=5):
        for j, v in enumerate([rule, sev, ent, CATEGORY[rule], n, BUSINESS_IMPACT[rule]], 1):
            c = ws.cell(row=i, column=j, value=v)
            c.border, c.alignment = border, Alignment(vertical="top", wrap_text=True)
            if j == 2:
                c.fill = SEV_FILL[sev]
    t = 5 + len(summary)
    ws.cell(row=t, column=1, value="Total").font = Font(bold=True)
    ws.cell(row=t, column=5, value=f"=SUM(E5:E{t - 1})").font = Font(bold=True)
    for col, w in zip("ABCDEF", (34, 10, 24, 28, 9, 80)):
        ws.column_dimensions[col].width = w

    ws2 = wb.create_sheet("Issue Register")
    heads = ["#", "Run", "Entity", "Business Key", "Category", "Rule Code", "Severity", "Description",
             "Business Impact", "Treatment"]
    for j, h in enumerate(heads, 1):
        c = ws2.cell(row=1, column=j, value=h)
        c.font, c.fill, c.border = hfont, hfill, border
    for i, (_, rid, ent, key, rule, sev, desc, treat) in enumerate(rows[:a.max_rows], start=2):
        for j, v in enumerate([i - 1, rid, ent, key, CATEGORY[rule], rule, sev, desc, BUSINESS_IMPACT[rule], treat], 1):
            c = ws2.cell(row=i, column=j, value=v)
            c.border, c.alignment = border, Alignment(vertical="top", wrap_text=True)
            if j == 7:
                c.fill = SEV_FILL[sev]
    for col, w in zip("ABCDEFGHIJ", (6, 6, 22, 14, 24, 30, 10, 48, 44, 48)):
        ws2.column_dimensions[col].width = w
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = f"A1:J{min(len(rows), a.max_rows) + 1}"

    a.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(a.out)
    print(f"wrote {a.out} ({len(rows)} exceptions, {len(summary)} rule groups)")


if __name__ == "__main__":
    main()
